"""Bounded attachment storage, inspection, scanning, and parsing.

Only server-generated attachment IDs are accepted after ingestion. Parsers run in a
short-lived subprocess so malformed customer files cannot monopolise the workflow
worker. All extracted text remains untrusted data.
"""

from __future__ import annotations

import hashlib
import io
import multiprocessing
import os
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from itertools import zip_longest
from multiprocessing.connection import Connection
from pathlib import Path, PurePath
from typing import Literal, Protocol

from pydantic import BaseModel, ConfigDict, Field


class AttachmentBoundaryError(ValueError):
    """A stable, review-safe attachment failure."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class AttachmentLimits:
    max_attachments: int = 10
    max_bytes_per_attachment: int = 10 * 1024 * 1024
    max_total_bytes: int = 25 * 1024 * 1024
    max_pdf_pages: int = 50
    max_pptx_slides: int = 100
    max_xlsx_sheets: int = 10
    max_xlsx_rows_per_sheet: int = 5_000
    max_xlsx_cells: int = 50_000
    max_extracted_chars: int = 100_000
    parser_timeout_seconds: float = 5.0
    max_zip_entries: int = 10_000
    max_zip_uncompressed_bytes: int = 100 * 1024 * 1024
    max_zip_compression_ratio: float = 100.0


class ParsedFragment(BaseModel):
    """One provenance-addressable fragment from an attachment."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    text: str = Field(min_length=1)
    page: int | None = Field(default=None, ge=1)
    slide: int | None = Field(default=None, ge=1)
    sheet: str | None = None
    row: int | None = Field(default=None, ge=1)
    cell: str | None = None


class ParsedAttachment(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    attachment_id: str
    media_type: str
    digest: str
    fragments: tuple[ParsedFragment, ...]
    normalized_text: str


class AttachmentScanner(Protocol):
    def scan(self, *, attachment_id: str, content: bytes, digest: str) -> Literal["clean"]: ...


class DeterministicAttachmentScanner:
    """Fixture scanner with digest/ID-addressed clean, infected, and timeout cases."""

    def __init__(
        self,
        *,
        infected: set[str] | None = None,
        timeouts: set[str] | None = None,
    ) -> None:
        self._infected = infected or set()
        self._timeouts = timeouts or set()

    def scan(self, *, attachment_id: str, content: bytes, digest: str) -> Literal["clean"]:
        del content
        keys = {attachment_id, digest}
        if keys & self._timeouts:
            raise AttachmentBoundaryError("MALWARE_SCAN_TIMEOUT", "Attachment scan timed out")
        if keys & self._infected:
            raise AttachmentBoundaryError("MALWARE_DETECTED", "Attachment failed malware screening")
        return "clean"


_ALLOWED_EXTENSIONS = {".pdf", ".pptx", ".xlsx"}
_MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_ATTACHMENT_ID = re.compile(r"^[a-f0-9]{64}$")


def normalize_untrusted_text(text: str) -> str:
    """Normalize for matching while the binary digest preserves original evidence."""

    return unicodedata.normalize("NFKC", text).replace("\x00", "")


def validate_original_filename(filename: str) -> str:
    """Reject path semantics even though the name is never used as a storage path."""

    if not filename or "\x00" in filename:
        raise AttachmentBoundaryError("UNSAFE_FILENAME", "Attachment filename is empty or invalid")
    if PurePath(filename).is_absolute() or "/" in filename or "\\" in filename or filename != Path(filename).name:
        raise AttachmentBoundaryError("UNSAFE_FILENAME", "Attachment filename contains path components")
    if filename in {".", ".."} or ".." in PurePath(filename).parts:
        raise AttachmentBoundaryError("UNSAFE_FILENAME", "Attachment filename contains traversal")
    extension = Path(filename).suffix.lower()
    if extension not in _ALLOWED_EXTENSIONS:
        raise AttachmentBoundaryError(
            "UNSUPPORTED_ATTACHMENT", f"Unsupported attachment extension: {extension or '<none>'}"
        )
    return extension


class AttachmentStore:
    """Content-addressed, per-message storage with no caller-selected path."""

    def __init__(self, root: Path) -> None:
        self.root = root.resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def attachment_id(mailbox_id: str, provider_message_id: str, ordinal: int, content: bytes) -> str:
        seed = b"\0".join(
            (
                mailbox_id.encode("utf-8"),
                provider_message_id.encode("utf-8"),
                str(ordinal).encode("ascii"),
                hashlib.sha256(content).digest(),
            )
        )
        return hashlib.sha256(seed).hexdigest()

    def put(self, *, mailbox_id: str, provider_message_id: str, ordinal: int, content: bytes) -> tuple[str, str]:
        attachment_id = self.attachment_id(mailbox_id, provider_message_id, ordinal, content)
        message_key = hashlib.sha256(f"{mailbox_id}\0{provider_message_id}".encode()).hexdigest()
        unresolved_directory = self.root / message_key
        if unresolved_directory.is_symlink():
            raise AttachmentBoundaryError("UNSAFE_STORAGE_REFERENCE", "Attachment directory cannot be a symlink")
        unresolved_directory.mkdir(mode=0o700, parents=True, exist_ok=True)
        directory = unresolved_directory.resolve()
        if self.root not in directory.parents:
            raise AttachmentBoundaryError("UNSAFE_STORAGE_REFERENCE", "Attachment directory escaped storage root")
        unresolved_path = directory / attachment_id
        if unresolved_path.is_symlink():
            raise AttachmentBoundaryError("UNSAFE_STORAGE_REFERENCE", "Attachment path cannot be a symlink")
        path = unresolved_path.resolve()
        if directory not in path.parents or path.is_symlink():
            raise AttachmentBoundaryError("UNSAFE_STORAGE_REFERENCE", "Attachment storage containment failed")
        if path.exists() and path.read_bytes() != content:
            raise AttachmentBoundaryError("ATTACHMENT_ID_COLLISION", "Attachment identity collision")
        if not path.exists():
            flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL
            descriptor = os.open(path, flags, 0o600)
            with os.fdopen(descriptor, "wb") as handle:
                handle.write(content)
        return attachment_id, f"{message_key}/{attachment_id}"

    def read(self, storage_reference: str, *, expected_attachment_id: str, expected_digest: str) -> bytes:
        if not _ATTACHMENT_ID.fullmatch(expected_attachment_id):
            raise AttachmentBoundaryError("UNSAFE_STORAGE_REFERENCE", "Invalid server attachment ID")
        relative = PurePath(storage_reference)
        if relative.is_absolute() or len(relative.parts) != 2 or relative.parts[-1] != expected_attachment_id:
            raise AttachmentBoundaryError("UNSAFE_STORAGE_REFERENCE", "Invalid attachment storage reference")
        unresolved_path = self.root / Path(*relative.parts)
        unresolved_parent = unresolved_path.parent
        if unresolved_parent.is_symlink() or unresolved_path.is_symlink():
            raise AttachmentBoundaryError("UNSAFE_STORAGE_REFERENCE", "Attachment storage cannot contain symlinks")
        path = unresolved_path.resolve()
        if self.root not in path.parents or not path.is_file():
            raise AttachmentBoundaryError("UNSAFE_STORAGE_REFERENCE", "Attachment storage containment failed")
        content = path.read_bytes()
        if hashlib.sha256(content).hexdigest() != expected_digest:
            raise AttachmentBoundaryError(
                "ATTACHMENT_DIGEST_MISMATCH", "Stored attachment digest does not match metadata"
            )
        return content


def _inspect_zip(content: bytes, *, extension: str, limits: AttachmentLimits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            names = {item.filename for item in infos}
            expected = "ppt/presentation.xml" if extension == ".pptx" else "xl/workbook.xml"
            if expected not in names or "[Content_Types].xml" not in names:
                raise AttachmentBoundaryError("WRONG_FILE_TYPE", "Office container is missing required entries")
            if len(infos) > limits.max_zip_entries:
                raise AttachmentBoundaryError("DECOMPRESSION_BOMB", "Office container has too many entries")
            total_compressed = sum(max(item.compress_size, 1) for item in infos)
            total_uncompressed = sum(item.file_size for item in infos)
            ratio = total_uncompressed / total_compressed
            if total_uncompressed > limits.max_zip_uncompressed_bytes or ratio > limits.max_zip_compression_ratio:
                raise AttachmentBoundaryError("DECOMPRESSION_BOMB", "Office container exceeds expansion limits")
            for item in infos:
                item_path = PurePath(item.filename)
                if item_path.is_absolute() or ".." in item_path.parts:
                    raise AttachmentBoundaryError("UNSAFE_CONTAINER_PATH", "Office container contains an unsafe entry")
    except AttachmentBoundaryError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise AttachmentBoundaryError("CORRUPT_ATTACHMENT", "Invalid Office container") from exc


def inspect_attachment(
    content: bytes,
    *,
    original_filename: str,
    declared_media_type: str | None,
    limits: AttachmentLimits,
) -> tuple[str, str]:
    extension = validate_original_filename(original_filename)
    if not content:
        raise AttachmentBoundaryError("EMPTY_ATTACHMENT", "Attachment is empty")
    if len(content) > limits.max_bytes_per_attachment:
        raise AttachmentBoundaryError("ATTACHMENT_TOO_LARGE", "Attachment exceeds the per-file limit")
    detected_media_type = _MEDIA_TYPES[extension]
    if declared_media_type and declared_media_type.lower() != detected_media_type:
        raise AttachmentBoundaryError("MIME_MISMATCH", "Declared and detected attachment types differ")
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise AttachmentBoundaryError("WRONG_FILE_TYPE", "PDF signature is missing")
    else:
        if not content.startswith(b"PK"):
            raise AttachmentBoundaryError("WRONG_FILE_TYPE", "Office ZIP signature is missing")
        _inspect_zip(content, extension=extension, limits=limits)
    return extension, detected_media_type


def _append_limited(parts: list[dict[str, object]], fragment: dict[str, object], total: int, maximum: int) -> int:
    text = normalize_untrusted_text(str(fragment.get("text", ""))).strip()
    if not text:
        return total
    new_total = total + len(text)
    if new_total > maximum:
        raise AttachmentBoundaryError("EXTRACTION_TOO_LARGE", "Extracted text exceeds the configured limit")
    fragment["text"] = text
    parts.append(fragment)
    return new_total


def _parse_pdf(content: bytes, limits: AttachmentLimits) -> list[dict[str, object]]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise AttachmentBoundaryError("ENCRYPTED_ATTACHMENT", "Encrypted PDFs are unsupported")
        if len(reader.pages) > limits.max_pdf_pages:
            raise AttachmentBoundaryError("PDF_PAGE_LIMIT", "PDF exceeds the page limit")
        result: list[dict[str, object]] = []
        total = 0
        for page_number, page in enumerate(reader.pages, 1):
            total = _append_limited(
                result,
                {"text": page.extract_text() or "", "page": page_number},
                total,
                limits.max_extracted_chars,
            )
        if not result:
            raise AttachmentBoundaryError("EMPTY_OR_SCANNED_PDF", "PDF contains no extractable text")
        return result
    except AttachmentBoundaryError:
        raise
    except Exception as exc:
        raise AttachmentBoundaryError("CORRUPT_ATTACHMENT", "PDF parsing failed") from exc


def _parse_pptx(content: bytes, limits: AttachmentLimits) -> list[dict[str, object]]:
    from pptx import Presentation

    try:
        presentation = Presentation(io.BytesIO(content))
        if len(presentation.slides) > limits.max_pptx_slides:
            raise AttachmentBoundaryError("PPTX_SLIDE_LIMIT", "PPTX exceeds the slide limit")
        result: list[dict[str, object]] = []
        total = 0
        for slide_number, slide in enumerate(presentation.slides, 1):
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    total = _append_limited(
                        result,
                        {"text": shape.text_frame.text, "slide": slide_number},
                        total,
                        limits.max_extracted_chars,
                    )
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text for cell in row.cells)
                        total = _append_limited(
                            result,
                            {"text": row_text, "slide": slide_number},
                            total,
                            limits.max_extracted_chars,
                        )
        if not result:
            raise AttachmentBoundaryError("EMPTY_ATTACHMENT_TEXT", "PPTX contains no extractable text")
        return result
    except AttachmentBoundaryError:
        raise
    except Exception as exc:
        raise AttachmentBoundaryError("CORRUPT_ATTACHMENT", "PPTX parsing failed") from exc


def _parse_xlsx(content: bytes, limits: AttachmentLimits) -> list[dict[str, object]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        formula_workbook = load_workbook(io.BytesIO(content), data_only=False, read_only=True)
        if len(workbook.worksheets) > limits.max_xlsx_sheets:
            raise AttachmentBoundaryError("XLSX_SHEET_LIMIT", "XLSX exceeds the sheet limit")
        result: list[dict[str, object]] = []
        total_chars = 0
        total_cells = 0
        for worksheet, formula_worksheet in zip(workbook.worksheets, formula_workbook.worksheets, strict=True):
            row_count = 0
            paired_rows = zip_longest(worksheet.iter_rows(), formula_worksheet.iter_rows(), fillvalue=())
            for row, formula_row in paired_rows:
                row_count += 1
                if row_count > limits.max_xlsx_rows_per_sheet:
                    raise AttachmentBoundaryError("XLSX_ROW_LIMIT", "XLSX exceeds the row limit")
                for cell, formula_cell in zip_longest(row, formula_row, fillvalue=None):
                    if (
                        formula_cell is not None
                        and formula_cell.data_type == "f"
                        and (cell is None or cell.value is None)
                    ):
                        raise AttachmentBoundaryError(
                            "XLSX_FORMULA_CACHE_MISSING",
                            "XLSX contains a formula without a cached value",
                        )
                    if cell is None:
                        continue
                    if cell.value is None:
                        continue
                    total_cells += 1
                    if total_cells > limits.max_xlsx_cells:
                        raise AttachmentBoundaryError("XLSX_CELL_LIMIT", "XLSX exceeds the cell limit")
                    total_chars = _append_limited(
                        result,
                        {
                            "text": str(cell.value),
                            "sheet": worksheet.title,
                            "row": cell.row,
                            "cell": cell.coordinate,
                        },
                        total_chars,
                        limits.max_extracted_chars,
                    )
        workbook.close()
        formula_workbook.close()
        if not result:
            raise AttachmentBoundaryError("EMPTY_ATTACHMENT_TEXT", "XLSX contains no cached values")
        return result
    except AttachmentBoundaryError:
        raise
    except Exception as exc:
        raise AttachmentBoundaryError("CORRUPT_ATTACHMENT", "XLSX parsing failed") from exc


def _parse_worker(
    connection: Connection,
    extension: str,
    content: bytes,
    limits: AttachmentLimits,
) -> None:
    try:
        parser = {".pdf": _parse_pdf, ".pptx": _parse_pptx, ".xlsx": _parse_xlsx}[extension]
        connection.send((True, parser(content, limits)))
    except AttachmentBoundaryError as exc:
        connection.send((False, exc.code, str(exc)))
    except BaseException:
        connection.send((False, "PARSER_FAILURE", "Attachment parser failed"))
    finally:
        connection.close()


def parse_in_subprocess(content: bytes, *, extension: str, limits: AttachmentLimits) -> tuple[ParsedFragment, ...]:
    context = multiprocessing.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    process = context.Process(target=_parse_worker, args=(child, extension, content, limits), daemon=True)
    process.start()
    child.close()
    try:
        if not parent.poll(limits.parser_timeout_seconds):
            process.terminate()
            process.join(timeout=1)
            raise AttachmentBoundaryError("PARSER_TIMEOUT", "Attachment parser exceeded its time limit")
        response = parent.recv()
    except EOFError as exc:
        raise AttachmentBoundaryError("PARSER_FAILURE", "Attachment parser exited unexpectedly") from exc
    finally:
        parent.close()
        if process.is_alive():
            process.terminate()
        process.join(timeout=1)
    if not response[0]:
        raise AttachmentBoundaryError(response[1], response[2])
    return tuple(ParsedFragment.model_validate(fragment) for fragment in response[1])


def scan_and_parse_attachment(
    *,
    attachment_id: str,
    content: bytes,
    original_filename: str,
    declared_media_type: str | None,
    expected_digest: str,
    scanner: AttachmentScanner,
    limits: AttachmentLimits | None = None,
) -> ParsedAttachment:
    limits = limits or AttachmentLimits()
    digest = hashlib.sha256(content).hexdigest()
    if digest != expected_digest:
        raise AttachmentBoundaryError("ATTACHMENT_DIGEST_MISMATCH", "Attachment digest does not match metadata")
    extension, media_type = inspect_attachment(
        content,
        original_filename=original_filename,
        declared_media_type=declared_media_type,
        limits=limits,
    )
    scanner.scan(attachment_id=attachment_id, content=content, digest=digest)
    fragments = parse_in_subprocess(content, extension=extension, limits=limits)
    normalized_text = "\n".join(fragment.text for fragment in fragments)
    return ParsedAttachment(
        attachment_id=attachment_id,
        media_type=media_type,
        digest=digest,
        fragments=fragments,
        normalized_text=normalized_text,
    )
